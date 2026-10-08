"""元バイト列を保存し、形式・取得履歴・日付の収録状況だけを検査する。"""

import codecs
import csv
import fcntl
import hashlib
import io
import json
import os
import re
import tempfile
import urllib.parse
import urllib.request
import zipfile
from collections import Counter
from datetime import date, datetime, timedelta
from email.message import Message
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from pypdf import PdfReader

JST = ZoneInfo("Asia/Tokyo")
FY_START = date(2025, 4, 1)
FY_END = date(2026, 3, 31)
MAX_BYTES = 100 * 1024 * 1024


def now():
    return datetime.now(JST).isoformat(timespec="seconds")


def date_range(start, end):
    if end < start:
        raise ValueError("期間の終了が開始より前")
    return [(start + timedelta(days=i)).isoformat() for i in range((end - start).days + 1)]


def detect_encoding(data):
    """BOM・厳密な復号で候補を記録する。曖昧なコードを断定しない。"""
    for bom, encoding, label in [
        (codecs.BOM_UTF32_LE, "utf-32", "UTF-32-LE"),
        (codecs.BOM_UTF32_BE, "utf-32", "UTF-32-BE"),
        (codecs.BOM_UTF8, "utf-8-sig", "UTF-8"),
        (codecs.BOM_UTF16_LE, "utf-16", "UTF-16-LE"),
        (codecs.BOM_UTF16_BE, "utf-16", "UTF-16-BE"),
    ]:
        if data.startswith(bom):
            return data.decode(encoding), {
                "encoding": encoding,
                "bom": label,
                "encoding_candidates": [encoding],
                "encoding_evidence": "BOMと全バイト列の厳密復号",
            }
    try:
        text = data.decode("utf-8")
        encoding = "ascii" if data.isascii() else "utf-8"
        return text, {
            "encoding": encoding,
            "bom": None,
            "encoding_candidates": [encoding],
            "encoding_evidence": "全バイト列の厳密復号。ASCIIの上位互換は判別不能",
        }
    except UnicodeDecodeError:
        pass
    candidates = []
    for encoding in ("cp932", "shift_jis", "euc_jp"):
        try:
            text = data.decode(encoding)
            if not any(ord(c) < 32 and c not in "\r\n\t" for c in text):
                candidates.append((encoding, text))
        except UnicodeDecodeError:
            pass
    if not candidates:
        raise ValueError("文字コードを厳密復号できない。rawを保存する前に要確認")
    # 候補間で同じ文字列の場合のみ検査に使用する。生バイトは変更しない。
    mapping_note = ""
    if len({text for _, text in candidates}) != 1:
        # CP932/SJISの同じバイトに対する既知のUnicode字形差のみ許容する。
        # この比較は判定根拠のためだけに使い、元ファイルや検査結果の列名は変換しない。
        variants = str.maketrans("〜‖−¢£¬", "～∥－￠￡￢")
        if {encoding for encoding, _ in candidates} != {"cp932", "shift_jis"} or len(
            {text.translate(variants) for _, text in candidates}
        ) != 1:
            raise ValueError("文字コード候補の復号結果が異なる。手動確認が必要")
        mapping_note = "。CP932/SJISのUnicode字形差あり（波ダッシュ等）"
    return candidates[0][1], {
        "encoding": candidates[0][0] if len(candidates) == 1 else "ambiguous",
        "bom": None,
        "encoding_candidates": [encoding for encoding, _ in candidates],
        "encoding_evidence": "候補の全バイト列を厳密復号" + mapping_note,
        "inspection_decoder": candidates[0][0],
    }


def parse_date(value):
    match = re.match(r"^(\d{4})[/-](\d{1,2})[/-](\d{1,2})(?:\s|$)", str(value).strip())
    return date(*map(int, match.groups())) if match else None


def inspect_csv(data, date_column=None, period_column=None):
    text, result = detect_encoding(data)
    if text.lstrip().lower().startswith(("<!doctype html", "<html")):
        raise ValueError("CSVではなくHTML応答")
    if not text.strip() or "\x00" in text:
        raise ValueError("空または不正なCSV応答")
    try:
        dialect = csv.Sniffer().sniff("\n".join(text.splitlines()[1:30]), delimiters=",\t;")
        delimiter = dialect.delimiter
    except csv.Error:
        # 更新日時等の1列プリアンブルを除き、最初の複数列の行を検査する。
        delimiter = next(
            (
                d
                for line in text.splitlines()[:30]
                for d in (",", "\t", ";")
                if len(next(csv.reader([line], delimiter=d))) > 1
            ),
            None,
        )
        if delimiter is None:
            raise ValueError("区切り文字を確認できない") from None
    reader = csv.reader(io.StringIO(text), delimiter=delimiter, strict=True)
    preamble = []
    columns = None
    header_row = None
    row_count = 0
    malformed = 0
    unparsed_dates = 0
    dates = Counter()
    periods = Counter()
    time_labels = set()
    time_index = None
    date_index = period_index = None
    for line_number, row in enumerate(reader, 1):
        if not row:
            continue
        if columns is None:
            if len(row) < 2:
                preamble.append(row)
                continue
            columns, header_row = row, line_number
            time_index = next((row.index(name) for name in ("時間帯", "時刻") if name in row), None)
            if date_column:
                if date_column not in row:
                    raise ValueError(f"対象日列がない: {date_column}")
                date_index = row.index(date_column)
            if period_column:
                if period_column not in row:
                    raise ValueError(f"コマ列がない: {period_column}")
                period_index = row.index(period_column)
            continue
        row_count += 1
        if len(row) != len(columns):
            malformed += 1
            continue
        if time_index is not None:
            time_labels.add(row[time_index])
        if date_index is not None:
            delivery_date = parse_date(row[date_index])
            if delivery_date is None:
                unparsed_dates += 1
                continue
            key = delivery_date.isoformat()
            dates[key] += 1
            if period_index is not None:
                periods[(key, row[period_index])] += 1
    if not columns:
        raise ValueError("CSV列名を確認できない")
    result.update(
        {
            "format": "csv",
            "delimiter": delimiter,
            "sheet_names": [],
            "header_row": header_row,
            "columns": columns,
            "preamble": preamble,
            "row_count": row_count,
            "malformed_rows": malformed,
            "unparsed_date_rows": unparsed_dates,
            "date_column": date_column,
            "observed_dates": sorted(dates),
            "period_start": min(dates) if dates else None,
            "period_end": max(dates) if dates else None,
            "observed_time_labels": sorted(time_labels),
        }
    )
    if period_column:
        result["period_audit"] = {
            "column": period_column,
            "duplicate_date_period_keys": sum(n - 1 for n in periods.values() if n > 1),
            "invalid_period_rows": sum(
                n for (_, p), n in periods.items() if p not in {str(i) for i in range(1, 49)}
            ),
            "incomplete_dates": [
                day
                for day in dates
                if {p for d, p in periods if d == day} != {str(i) for i in range(1, 49)}
            ],
        }
    return result


def inspect_bytes(data, filename, mime_type=None, date_column=None, period_column=None):
    if not data:
        raise ValueError("空応答")
    if len(data) > MAX_BYTES:
        raise ValueError("ファイルサイズ上限超過")
    ext = Path(filename).suffix.lower()
    base = {"extension": ext, "http_mime_type": mime_type}
    if data.startswith(b"%PDF-") and ext == ".pdf":
        reader = PdfReader(io.BytesIO(data))
        base.update(
            {
                "format": "pdf",
                "page_count": len(reader.pages),
                "encoding": "not_applicable",
                "bom": None,
                "delimiter": None,
                "sheet_names": [],
                "columns": [],
                "period_start": None,
                "period_end": None,
                "period_evidence": "PDFの対象期間は手動確認が必要",
            }
        )
    elif zipfile.is_zipfile(io.BytesIO(data)) and ext in {".zip", ".xlsx"}:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(i.file_size for i in archive.infolist()) > MAX_BYTES:
                raise ValueError("展開後サイズ上限超過")
            if ext == ".xlsx":
                workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
                base.update(
                    {
                        "format": "xlsx",
                        "sheet_names": workbook.sheetnames,
                        "columns": {
                            s.title: [
                                str(v) if v is not None else None
                                for v in next(s.iter_rows(values_only=True), [])
                            ]
                            for s in workbook.worksheets
                        },
                        "encoding": "not_applicable",
                        "bom": None,
                        "delimiter": None,
                        "period_start": None,
                        "period_end": None,
                    }
                )
                workbook.close()
            else:
                members = {}
                for item in archive.infolist():
                    if item.filename.lower().endswith(".csv"):
                        members[item.filename] = inspect_csv(archive.read(item), date_column)
                base.update(
                    {
                        "format": "zip",
                        "members": members,
                        "member_names": archive.namelist(),
                        "encoding": "per_member",
                        "bom": None,
                        "delimiter": None,
                        "sheet_names": [],
                        "columns": [],
                        "period_start": None,
                        "period_end": None,
                    }
                )
    elif ext == ".csv":
        base.update(inspect_csv(data, date_column, period_column))
    else:
        raise ValueError(f"形式と拡張子を確認できない: {ext}")
    return base


def save_exclusive(path, data):
    """同一内容でも既存rawは開かない。競合時にも上書きしない。"""
    path = Path(path)
    if path.exists() or path.is_symlink():
        raise FileExistsError(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(dir=path.parent, delete=False) as handle:
        temporary = Path(handle.name)
        try:
            handle.write(data)
            handle.flush()
            os.fsync(handle.fileno())
            os.link(temporary, path)
        finally:
            temporary.unlink(missing_ok=True)


def append_manifest(path, record):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a", encoding="utf-8") as handle:
        fcntl.flock(handle, fcntl.LOCK_EX)
        handle.write(json.dumps(record, ensure_ascii=False) + "\n")
        handle.flush()
        fcntl.flock(handle, fcntl.LOCK_UN)


def read_manifest(path):
    path = Path(path)
    if not path.exists():
        return []
    with path.open(encoding="utf-8") as handle:
        fcntl.flock(handle, fcntl.LOCK_SH)
        records = [json.loads(line) for line in handle if line.strip()]
        fcntl.flock(handle, fcntl.LOCK_UN)
    return records


def original_filename(disposition, fallback):
    message = Message()
    message["Content-Disposition"] = disposition or ""
    return Path(urllib.parse.unquote(message.get_filename() or fallback)).name


def fetch(url, fields=None, referer=None, opener=None, extra_headers=None):
    headers = {"User-Agent": "jepx-market-power-analysis/0.1", "Accept-Encoding": "identity"}
    if referer:
        headers["Referer"] = referer
    headers.update(extra_headers or {})
    payload = urllib.parse.urlencode(fields, doseq=True).encode() if fields is not None else None
    request = urllib.request.Request(url, data=payload, headers=headers)
    client = opener or urllib.request.build_opener()
    with client.open(request, timeout=45) as response:
        data = response.read(MAX_BYTES + 1)
        if len(data) > MAX_BYTES:
            raise ValueError("HTTP応答サイズ上限超過")
        return data, dict(response.headers), response.url


def record_download(
    root,
    source_id,
    source_name,
    url,
    data,
    headers,
    final_url,
    filename,
    requested_start=None,
    requested_end=None,
    date_column=None,
    period_column=None,
    fields=None,
    notes="",
    phase="sample",
    variant="default",
):
    root = Path(root)
    # 保存前にHTML/空/不正形式を排除。拡張子やHTTP宣言だけでは成功扱いしない。
    lower_headers = {k.lower(): v for k, v in headers.items()}
    inspection = inspect_bytes(
        data, filename, lower_headers.get("content-type"), date_column, period_column
    )
    path = root / "data" / "raw" / source_id / filename
    if Path(filename).name != filename or filename in {"", ".", ".."}:
        raise ValueError("不正な保存ファイル名")
    if not re.fullmatch(r"[a-z][a-z0-9_]+", source_id):
        raise ValueError("不正なソースID")
    if not path.resolve().is_relative_to((root.resolve() / "data" / "raw")):
        raise ValueError("raw保存先がリポジトリ外へ解決される")
    save_exclusive(path, data)
    record = {
        "source_id": source_id,
        "source_name": source_name,
        "source_url": url,
        "final_url": final_url,
        "http_method": "POST" if fields is not None else "GET",
        "request_parameters": {
            k: v
            for k, v in (fields or {}).items()
            if not any(part in k.lower() for part in ("csrf", "token", "session", "downloadkey"))
        },
        "local_path": str(path.relative_to(root)),
        "original_filename": original_filename(lower_headers.get("content-disposition"), filename),
        "requested_period_start": requested_start,
        "requested_period_end": requested_end,
        "period_start": inspection.get("period_start"),
        "period_end": inspection.get("period_end"),
        "downloaded_at": now(),
        "file_size": len(data),
        "sha256": hashlib.sha256(data).hexdigest(),
        "status": "downloaded",
        "phase": phase,
        "variant": variant,
        "notes": notes,
        "inspection": inspection,
    }
    if requested_start and requested_end and date_column and inspection.get("format") == "csv":
        expected = date_range(
            date.fromisoformat(requested_start), date.fromisoformat(requested_end)
        )
        observed = set(inspection["observed_dates"])
        record["missing_dates"] = sorted(set(expected) - observed)
        record["unexpected_dates"] = sorted(observed - set(expected))
    append_manifest(root / "metadata" / "acquisition_manifest.jsonl", record)
    return record
