"""保存した少数公式サンプルを再検査。ネットワーク取得や年度巡回は行わない。"""

import argparse
import hashlib
import json
import sys
from pathlib import Path

import openpyxl
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from src.fixed_output_feasibility import (  # noqa: E402
    decode_csv,
    inspect_half_hours,
    limited_supply_register,
)


def sha(path):
    with path.open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest()


def run(root, output):
    output = output.resolve()
    if output.exists():
        raise ValueError("既存出力を上書きしない")
    output.mkdir(parents=True)
    report = root / "reports/18_fixed_output_feasibility/qa"
    report.mkdir(parents=True, exist_ok=True)
    reference_manifest = root / "metadata/research/R_feasibility_reference_manifest_20261026.json"
    refs = json.loads(reference_manifest.read_text())
    for r in refs:
        if "path" in r and sha(root / r["path"]) != r["sha256"]:
            raise ValueError("取得原本SHA不一致")
    sample = root / "data/interim/research_reference/20261026_R_feasibility_v1"
    audit = {"samples": {}, "input_references": refs, "real_network_lp_runs": 0}
    paths = []
    for name, header_row, time_columns in [
        ("hokuriku_fit_202504_06.csv", 0, ["DATE_TIME"]),
        ("tokyo_actual_202504.csv", 1, ["DATE", "TIME"]),
        ("shikoku_forecast_today.csv", 2, ["日付", "時間帯_自"]),
    ]:
        rows, encoding = decode_csv((sample / name).read_bytes())
        header = rows[header_row]
        data = rows[header_row + 1 :]
        if any(len(r) != len(header) for r in data):
            raise ValueError("CSV列数不一致")
        frame = pd.DataFrame(data, columns=header)
        if name.startswith("shikoku"):
            times = pd.to_datetime(frame["日付"], format="%Y%m%d") + pd.to_timedelta(
                frame["時間帯_自"].map(lambda x: x + ":00")
            )
            numeric = header[4:]
        elif name.startswith("tokyo"):
            times = frame["DATE"] + " " + frame["TIME"]
            numeric = header[2:]
        else:
            times = frame["DATE_TIME"]
            numeric = header[1:]
        enriched, qa = inspect_half_hours(frame, times, numeric)
        if not enriched[header].equals(frame):
            raise ValueError("サンプル元セル変更")
        destination = output / (name.removesuffix(".csv") + "_inspected.parquet")
        enriched.to_parquet(destination, index=False)
        if not pd.read_parquet(destination).equals(enriched):
            raise ValueError("サンプル再読込不一致")
        qa.update(
            encoding=encoding,
            utf8_bom=(sample / name).read_bytes().startswith(b"\xef\xbb\xbf"),
            columns=header,
            header_rows=rows[: header_row + 1],
            prediction_issue_timestamp_in_data=False,
        )
        if name.startswith("shikoku"):
            qa["file_update_fields"] = rows[1]
        audit["samples"][name] = qa
        paths.append(destination)
    # Excelは全セルの座標・型・式を別表で保存。日量や空の31日列も落とさない。
    excel = sample / "chubu_fit_202504.xlsx"
    book = openpyxl.load_workbook(excel, data_only=False)
    cells, worksheets = [], {}
    for sheet in book:
        for row in sheet:
            for cell in row:
                cells.append(
                    {
                        "sheet": sheet.title,
                        "coordinate": cell.coordinate,
                        "raw_value": None if cell.value is None else str(cell.value),
                        "cell_type": cell.data_type,
                    }
                )
        dates = [
            (col, sheet.cell(4, col).value)
            for col in range(2, sheet.max_column + 1)
            if hasattr(sheet.cell(4, col).value, "date")
        ]
        frame = pd.DataFrame(
            [
                {
                    "sheet": sheet.title,
                    "date_raw": str(day),
                    "time_range_raw": sheet.cell(p + 4, 1).value,
                    "value_raw": ""
                    if sheet.cell(p + 4, col).value is None
                    else str(sheet.cell(p + 4, col).value),
                    "source_cell": sheet.cell(p + 4, col).coordinate,
                }
                for col, day in dates
                for p in range(1, 49)
            ]
        )
        times = [day + pd.Timedelta(minutes=30 * (p - 1)) for _, day in dates for p in range(1, 49)]
        _, qa = inspect_half_hours(frame, times, ["value_raw"])
        qa["unit_cell"] = str(sheet["AC3"].value)
        qa["forecast_stage_evidence"] = sheet.title
        qa["daily_total_mismatches"] = sum(
            sum(sheet.cell(p + 4, col).value for p in range(1, 49)) != sheet.cell(53, col).value
            for col, _ in dates
        )
        worksheets[sheet.title] = qa
    destination = output / "chubu_workbook_all_cells.parquet"
    pd.DataFrame(cells).to_parquet(destination, index=False)
    paths.append(destination)
    audit["samples"]["chubu_fit_202504.xlsx"] = worksheets
    inventory_path = root / "reports/17_net_demand_network/qa/fixed_supply_overlap_inventory.csv"
    business_path = root / (
        "data/processed/comparison/20261016_decision_preparation_v2/"
        "observed_unit_business_evidence_FY2025.parquet"
    )
    inventory = pd.read_csv(inventory_path, keep_default_na=False)
    business = pd.read_parquet(business_path)
    merged = inventory.merge(
        business,
        on="asset_key_candidate",
        how="left",
        validate="one_to_one",
        indicator="evidence_join",
        suffixes=("", "_business"),
    )
    if len(merged) != 491 or not merged["evidence_join"].eq("both").all():
        raise ValueError("491ラベルの根拠照合不一致")
    ref = pd.to_numeric(merged["diagnostic_mapping__capacity_reference_raw_kw"], errors="raise")
    merged["reference_capacity_positive_unique"] = ref.gt(0) & merged[
        "diagnostic_mapping__research_use_status"
    ].eq("use_diagnostic")
    # 既存状態の実値を確認。単に容量の正値だけで保留を解除しない。
    audit["research_use_status_counts"] = (
        merged["diagnostic_mapping__research_use_status"].value_counts().to_dict()
    )
    register = limited_supply_register(merged)
    columns = [
        "asset_key_candidate",
        "area",
        "plant_name",
        "unit_name",
        "発電方式・燃種",
        "generation_kwh_30min_sum",
        "diagnostic_mapping__capacity_reference_raw_kw",
        "diagnostic_mapping__research_use_status",
        "pv_wind_disjoint_by_raw_fuel",
        "limited_scope_status",
        "unknown_capacity_assumed_zero",
        "actual_availability_verified",
    ]
    register[columns].to_csv(report / "limited_supply_asset_register.csv", index=False)
    unknown = merged[merged["発電方式・燃種"].isin(["", "その他"])].copy()
    unit_path = root / (
        "data/raw/hjks/units/hjks_units__hjks_units_FY2025__sha256_7d4c3a8633fa.csv"
    )
    raw_rows, _ = decode_csv(unit_path.read_bytes())
    reviews = []
    for r in unknown.to_dict("records"):
        row_id = r["diagnostic_mapping__target_snapshot_row_id"]
        row = raw_rows[int(row_id.rsplit(":", 1)[1]) - 1] if row_id else None
        reviews.append(
            {
                "asset_key_candidate": r["asset_key_candidate"],
                "area": r["area"],
                "plant_name_raw": r["plant_name"],
                "unit_name_raw": r["unit_name"],
                "generation_fuel_raw": r["発電方式・燃種"],
                "hjks_generation_form_raw": row[4] if row else None,
                "hjks_row_id": row_id,
                "business_candidate_rows": r["hjks_candidate_rows_json"],
                "generation_kwh_observed": r["generation_kwh_30min_sum"],
                "reference_capacity_kw": r["diagnostic_mapping__capacity_reference_raw_kw"],
                "saved_evidence_pv_wind_overlap_status": "unresolved",
                "reason": (
                    "保存台帳の発電形式も空欄/その他。名称や法人名だけで非PV/風力と認定しない。"
                    if row
                    else "点検済み設備対応なし。候補台帳の名称から燃種を確定しない。"
                ),
                "source_file": str(unit_path.relative_to(root)),
                "source_sha256": sha(unit_path),
            }
        )
    pd.DataFrame(reviews).to_csv(report / "thirteen_asset_overlap_review.csv", index=False)
    audit["overlap_review"] = {
        "labels": len(reviews),
        "resolved": 0,
        "unresolved": len(reviews),
        "generation_kwh": str(unknown["generation_kwh_30min_sum"].sum()),
    }
    audit["limited_register_counts"] = register["limited_scope_status"].value_counts().to_dict()
    audit["business_join"] = {"left": 491, "matched": 491, "unmatched": 0, "duplicated": 0}
    # 切替前の同一時刻だけで旧9条件の既知部分和を比較。未知をゼロと認定しない。
    legacy = root / (
        "data/interim/occto/net_demand_network/20261026_diagnostic_v1/"
        "network_time_readiness_FY2025.parquet"
    )
    times = pd.read_parquet(legacy)
    candidate_times = times[times["network_structure_status"].eq("compiled_saved_plan_legacy")]
    regional_path = root / (
        "data/interim/occto/saved_plan_supply/20261024_diagnostic_v1/regional_all_suppliers.parquet"
    )
    regional = pd.read_parquet(regional_path)
    regional = regional[regional["area"].ne("沖縄")]
    selected = regional.merge(
        candidate_times[["delivery_date", "period"]],
        on=["delivery_date", "period"],
        how="left",
        validate="many_to_one",
        indicator="period_join",
    )
    audit["network_period_join"] = selected["period_join"].value_counts().to_dict()
    selected = selected[selected["period_join"].eq("both")]
    summaries = []
    for (scenario, area), group in selected.groupby(["supply_scenario", "area"]):
        x = group["adjusted_upper_partial_mw"].dropna()
        summaries.append(
            {
                "scenario": scenario,
                "area": area,
                "candidate_slots": len(group),
                "known_partial_min_mw": str(x.min()),
                "known_partial_max_mw": str(x.max()),
                "record_held_cells": int(group["record_held_count"].sum()),
                "capacity_held_cells": int(group["capacity_held_count"].sum()),
                "role": "既存既知部分和の入力範囲。新しい非重複集合の供給力ではない",
            }
        )
    pd.DataFrame(summaries).to_csv(report / "precutover_known_supply_scope.csv", index=False)
    audit["candidate_period"] = {
        "slots": len(candidate_times),
        "first": str(candidate_times["delivery_date"].min()),
        "last": str(candidate_times["delivery_date"].max()),
        "not_annual": True,
    }
    inputs = [inventory_path, business_path, unit_path, legacy, regional_path]
    audit["preserved_input_hashes"] = [
        {"path": str(p.relative_to(root)), "sha256": sha(p)} for p in inputs
    ]
    audit["outputs"] = [{"path": str(p.relative_to(root)), "sha256": sha(p)} for p in paths]
    audit["reference_count"] = len(refs)
    (report / "sample_audit.json").write_text(
        json.dumps(audit, ensure_ascii=False, indent=2) + "\n"
    )
    return audit


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    run(Path(__file__).resolve().parents[1], args.output)
